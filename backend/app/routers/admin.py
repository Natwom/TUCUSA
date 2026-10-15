from datetime import datetime, timedelta, timezone
from io import BytesIO
from itertools import groupby
from typing import Optional
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.database import get_db
from app import models, schemas
from app.auth import require_admin
from app.core.security_events import log_event

router = APIRouter(prefix="/api/admin", tags=["Admin"])


@router.get("/stats", response_model=schemas.AdminStats)
def get_stats(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    total_students = db.query(models.User).filter(models.User.role == "student").count()
    total_votes = db.query(models.Vote).count()
    active_elections = db.query(models.Election).filter(models.Election.status == "active").count()
    pending = (
        db.query(models.User)
        .filter(models.User.role == "student", models.User.is_approved == False)  # noqa: E712
        .count()
    )

    return schemas.AdminStats(
        total_students=total_students,
        total_votes_cast=total_votes,
        active_elections=active_elections,
        pending_approvals=pending,
    )


@router.get("/security-events", response_model=list[schemas.SecurityEventOut])
def list_security_events(
    severity: Optional[str] = None,
    event_type: Optional[str] = None,
    limit: int = Query(100, ge=1, le=500),
    offset: int = Query(0, ge=0),
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    q = db.query(models.SecurityEvent)
    if severity:
        q = q.filter(models.SecurityEvent.severity == severity)
    if event_type:
        q = q.filter(models.SecurityEvent.event_type == event_type)
    return q.order_by(models.SecurityEvent.created_at.desc()).offset(offset).limit(limit).all()


@router.get("/security-events/summary")
def security_summary(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    since = datetime.now(timezone.utc) - timedelta(hours=24)
    rows = (
        db.query(models.SecurityEvent.event_type, models.SecurityEvent.severity, func.count(models.SecurityEvent.id))
        .filter(models.SecurityEvent.created_at >= since)
        .group_by(models.SecurityEvent.event_type, models.SecurityEvent.severity)
        .all()
    )
    by_type, critical, warning = {}, 0, 0
    for event_type, severity, n in rows:
        by_type[event_type] = by_type.get(event_type, 0) + n
        if severity == "critical":
            critical += n
        elif severity == "warning":
            warning += n
    return {"last_24h": {"total": sum(by_type.values()), "critical": critical, "warning": warning, "by_type": by_type}}


@router.get("/integrity-check")
def integrity_check(
    request: Request,
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    """Compares each candidate's cached votes_count with the real number of rows in `votes`."""
    actual = dict(
        db.query(models.Vote.candidate_id, func.count(models.Vote.id)).group_by(models.Vote.candidate_id).all()
    )
    mismatches = []
    for c in db.query(models.Candidate).all():
        real = actual.get(c.id, 0)
        if real != (c.votes_count or 0):
            mismatches.append({"candidate_id": c.id, "name": c.name, "stored": c.votes_count, "actual": real})
    if mismatches:
        log_event(
            "integrity_mismatch", "critical", request, email=admin.email, user_id=admin.id,
            detail=f"{len(mismatches)} candidate(s) have vote counts that do not match the votes table",
        )
    return {"ok": not mismatches, "mismatches": mismatches}


# ---------------------------------------------------------------------------
# Voter register export (Word) for clerks
# ---------------------------------------------------------------------------

# (column title, width in cm). Landscape A4 with 1.5cm margins = 26.7cm usable.
REGISTER_COLUMNS = [
    ("No.", 1.0),
    ("Full Name", 4.5),
    ("Admission No.", 2.8),
    ("Course", 4.0),
    ("Year", 1.0),
    ("Constituency", 3.0),
    ("Email", 4.2),
    ("Phone", 2.6),
    ("Signature", 3.5),
]


def _query_voters(db: Session, constituency: Optional[str], status_filter: str):
    q = db.query(models.User).filter(models.User.role == "student")

    if constituency and constituency.strip().lower() != "all":
        q = q.filter(func.lower(models.User.constituency) == constituency.strip().lower())

    if status_filter == "pending":
        q = q.filter(models.User.is_approved == False, models.User.is_active == True)  # noqa: E712
    elif status_filter == "approved":
        q = q.filter(models.User.is_approved == True)  # noqa: E712
    elif status_filter == "rejected":
        q = q.filter(models.User.is_active == False)  # noqa: E712

    voters = q.all()
    voters.sort(key=lambda v: ((v.constituency or "").lower(), (v.full_name or "").lower()))
    return voters


def _repeat_header_row(row):
    """Repeat the table header row at the top of every page."""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _shade_cell(cell, hex_fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _write_cell(cell, text: str, width_cm: float, bold: bool = False):
    cell.width = Cm(width_cm)
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(9)
    run.bold = bold


def _add_voter_table(doc: Document, voters):
    table = doc.add_table(rows=1, cols=len(REGISTER_COLUMNS))
    table.style = "Table Grid"
    table.autofit = False

    header = table.rows[0]
    _repeat_header_row(header)
    for i, (title, width) in enumerate(REGISTER_COLUMNS):
        _write_cell(header.cells[i], title, width, bold=True)
        _shade_cell(header.cells[i], "D9E2F3")

    for n, v in enumerate(voters, start=1):
        row = table.add_row()
        row.height = Cm(0.9)  # room for a signature
        values = [
            str(n),
            v.full_name or "",
            v.admission_number or "",
            v.course or "",
            str(v.year_of_study or ""),
            v.constituency or "",
            v.email or "",
            v.phone or "",
            "",  # signature column left blank
        ]
        for i, val in enumerate(values):
            _write_cell(row.cells[i], val, REGISTER_COLUMNS[i][1])


def _build_register_docx(voters, constituency: Optional[str], status_filter: str) -> BytesIO:
    doc = Document()

    # Landscape A4, narrow margins
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(1.5)

    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    # Title block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TUCUSA VOTERS REGISTER")
    run.bold = True
    run.font.size = Pt(16)

    is_all = not constituency or constituency.strip().lower() == "all"
    scope = "All constituencies" if is_all else f"Constituency: {constituency}"
    status_label = {
        "all": "All statuses", "approved": "Approved voters only",
        "pending": "Pending voters only", "rejected": "Rejected voters only",
    }[status_filter]

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run(f"{scope}  |  {status_label}  |  Generated: {datetime.now().strftime('%d %B %Y, %H:%M')}")

    total = doc.add_paragraph()
    total.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total_run = total.add_run(f"Total voters: {len(voters)}")
    total_run.bold = True
    total_run.font.size = Pt(12)

    # One section per constituency (a single one when a constituency is chosen)
    groups = [(name, list(items)) for name, items in groupby(voters, key=lambda v: v.constituency)]
    for idx, (name, items) in enumerate(groups):
        if idx > 0:
            doc.add_page_break()
        heading = doc.add_paragraph()
        h_run = heading.add_run(f"Constituency: {name}  ({len(items)} voter{'s' if len(items) != 1 else ''})")
        h_run.bold = True
        h_run.font.size = Pt(12)
        _add_voter_table(doc, items)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@router.get("/voters/export-docx")
def export_voters_docx(
    request: Request,
    constituency: Optional[str] = Query(None, max_length=100),
    status: str = Query("all", pattern="^(all|pending|approved|rejected)$"),
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    voters = _query_voters(db, constituency, status)
    if not voters:
        raise HTTPException(status_code=404, detail="No voters match the selected filters")

    buf = _build_register_docx(voters, constituency, status)

    log_event(
        "voter_list_exported", "info", request, email=admin.email, user_id=admin.id,
        detail=f"constituency={constituency or 'all'} status={status} count={len(voters)}",
    )

    slug = re.sub(r"[^A-Za-z0-9]+", "_", constituency).strip("_") if constituency and constituency.lower() != "all" else "all"
    filename = f"tucusa_voters_{slug}_{datetime.now().strftime('%Y-%m-%d')}.docx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Voter list export (Excel)
# ---------------------------------------------------------------------------

# (column title, width in characters)
EXCEL_COLUMNS = [
    ("No.", 6),
    ("Full Name", 28),
    ("Admission No.", 18),
    ("Course", 30),
    ("Year", 7),
    ("Constituency", 20),
    ("Email", 32),
    ("Phone", 16),
    ("Voter ID", 18),
    ("Status", 12),
    ("Registered", 18),
]


def _voter_status(v) -> str:
    if v.is_approved:
        return "Approved"
    if not v.is_active:
        return "Rejected"
    return "Pending"


def _safe_sheet_title(name: str, used: set) -> str:
    """Excel sheet names: max 31 chars, no []:*?/\\ and must be unique."""
    base = re.sub(r"[\[\]\:\*\?\/\\]", " ", name or "Unknown").strip()[:31] or "Unknown"
    title, i = base, 2
    while title.lower() in used:
        suffix = f" ({i})"
        title = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(title.lower())
    return title


def _fill_voter_sheet(ws, voters):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")

    ws.append([title for title, _ in EXCEL_COLUMNS])
    for col_idx, (_, width) in enumerate(EXCEL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for n, v in enumerate(voters, start=1):
        registered = v.created_at.strftime("%Y-%m-%d %H:%M") if v.created_at else ""
        ws.append([
            n,
            v.full_name or "",
            v.admission_number or "",
            v.course or "",
            v.year_of_study or "",
            v.constituency or "",
            v.email or "",
            v.phone or "",
            v.unique_voter_id or "",
            _voter_status(v),
            registered,
        ])

    # Students type their own details, so never let a value be treated as a formula
    # (e.g. a name starting with "="). Force every text cell to plain string.
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _build_voters_xlsx(voters, constituency: Optional[str]) -> BytesIO:
    wb = Workbook()
    used: set = set()

    ws_all = wb.active
    ws_all.title = _safe_sheet_title("All Voters", used)
    _fill_voter_sheet(ws_all, voters)

    # When exporting several constituencies, add one sheet per constituency too
    is_all = not constituency or constituency.strip().lower() == "all"
    if is_all:
        groups = [(name, list(items)) for name, items in groupby(voters, key=lambda v: v.constituency)]
        if len(groups) > 1:
            for name, items in groups:
                ws = wb.create_sheet(_safe_sheet_title(name, used))
                _fill_voter_sheet(ws, items)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.get("/voters/export-xlsx")
def export_voters_xlsx(
    request: Request,
    constituency: Optional[str] = Query(None, max_length=100),
    status: str = Query("all", pattern="^(all|pending|approved|rejected)$"),
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    voters = _query_voters(db, constituency, status)
    if not voters:
        raise HTTPException(status_code=404, detail="No voters match the selected filters")

    buf = _build_voters_xlsx(voters, constituency)

    log_event(
        "voter_list_exported", "info", request, email=admin.email, user_id=admin.id,
        detail=f"format=xlsx constituency={constituency or 'all'} status={status} count={len(voters)}",
    )

    slug = re.sub(r"[^A-Za-z0-9]+", "_", constituency).strip("_") if constituency and constituency.lower() != "all" else "all"
    filename = f"tucusa_voters_{slug}_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )